from openpyxl import Workbook
import os

def build_excel(summary, month):

    os.makedirs("reports", exist_ok=True)   # ✅ ensure folder exists

    file_path = f"reports/report_{month}.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"

    ws.append(["Month", month])
    ws.append([])
    ws.append(["Total Revenue", summary["total_revenue"]])
    ws.append(["Total Expenses", summary["total_expenses"]])
    ws.append(["Net Profit", summary["net_profit"]])

    ws.append([])
    ws.append(["Category", "Amount"])

    for c in summary["category_data"]:
        ws.append([c["name"], c["value"]])

    wb.save(file_path)

    print("Saved report:", file_path)   # ✅ debug line

    return file_path


# ================= FULL FORMATTED REPORT =================

from io import BytesIO
from datetime import datetime
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# stored account name -> label shown in the report
SOURCE_LABELS = [
    ("CASH", "Cash"),
    ("NAYEEM PRIMARY", "Nayeem Primary"),
    ("NAYEEM CURRENT", "Nayeem Current"),
    ("SUBHAN KHAN", "Subhan Khan"),
    ("MAHABOOB BI", "Mahaboob bi"),
    ("MUBEENA", "Mubeena"),
]

GOLD = PatternFill("solid", fgColor="F1C232")
LIGHT = PatternFill("solid", fgColor="F3F3F3")
GREY = Side(style="thin", color="BFBFBF")
BORDER = Border(left=GREY, right=GREY, top=GREY, bottom=GREY)
RUPEE = '"₹"#,##0.00'


def build_full_report_workbook(data, month):
    """
    data = {
      "total_revenue": float, "total_expenses": float, "net_profit": float,
      "sources": {"CASH": 0, "NAYEEM PRIMARY": 0, ...},
      "categories": [{"name": "Mess", "value": 123}, ...]
    }
    """
    try:
        month_label = datetime.strptime(month, "%Y-%m").strftime("%B %Y")
    except ValueError:
        month_label = month

    wb = Workbook()
    ws = wb.active
    ws.title = "REVENUE"
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 24

    row = 1

    def band(text, size=12):
        nonlocal row
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        c = ws.cell(row=row, column=1, value=text)
        c.font = Font(bold=True, size=size)
        c.fill = GOLD
        c.alignment = Alignment(horizontal="center", vertical="center")
        for col in (1, 2):
            ws.cell(row=row, column=col).fill = GOLD
            ws.cell(row=row, column=col).border = BORDER
        ws.row_dimensions[row].height = 24
        row += 1

    def line(label, value, bold=False, color=None, fill=None):
        nonlocal row
        a = ws.cell(row=row, column=1, value=label)
        b = ws.cell(row=row, column=2, value=value)
        b.number_format = RUPEE
        b.alignment = Alignment(horizontal="right")
        font = Font(bold=bold, color=color) if color else Font(bold=bold)
        a.font = font
        b.font = font
        for c in (a, b):
            c.border = BORDER
            if fill:
                c.fill = fill
        ws.row_dimensions[row].height = 20
        row += 1

    def gap():
        nonlocal row
        row += 1

    # ---- title ----
    band("Khan's Hostel - Monthly Report", size=14)
    band(month_label)
    gap()

    # ---- summary ----
    line("Total Revenue", data["total_revenue"], bold=True)
    line("Total Expenses", data["total_expenses"], bold=True)
    profit = data["net_profit"]
    line("Net Profit", profit, bold=True,
         color="188038" if profit >= 0 else "D93025", fill=LIGHT)
    gap()

    # ---- revenue by source ----
    band("Revenue by Source")
    for key, label in SOURCE_LABELS:
        line(label, data["sources"].get(key, 0))
    gap()

    # ---- expenses by category ----
    band("Expenses by Category")
    hdr_a = ws.cell(row=row, column=1, value="Category")
    hdr_b = ws.cell(row=row, column=2, value="Amount")
    for c in (hdr_a, hdr_b):
        c.font = Font(bold=True)
        c.fill = GOLD
        c.border = BORDER
        c.alignment = Alignment(horizontal="center")
    row += 1

    for cat in sorted(data["categories"], key=lambda c: str(c["name"]).lower()):
        line(cat["name"], cat["value"])

    line("Total", data["total_expenses"], bold=True, fill=LIGHT)

    ws.sheet_view.showGridLines = False
    return wb


def build_full_report_bytes(data, month):
    stream = BytesIO()
    build_full_report_workbook(data, month).save(stream)
    stream.seek(0)
    return stream


def build_full_report_file(data, month):
    os.makedirs("reports", exist_ok=True)
    file_path = f"reports/report_{month}.xlsx"
    build_full_report_workbook(data, month).save(file_path)
    return file_path